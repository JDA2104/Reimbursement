"""Pembuat form rekap Excel.

Meniru layout Daftar Nominatif Entertainment yang dipakai perusahaan: blok
identitas karyawan di atas, tabel dua tingkat di tengah, lalu blok total dan
kolom tanda tangan di bawah. Satu file berisi dua sheet - Reimbursement dan
Kartu Kredit - dengan struktur yang identik.

Foto struk sengaja TIDAK disisipkan ke dalam form, supaya layout dokumen resmi
tidak berubah. Foto dikirim terpisah dalam .zip, dinamai sesuai nomor barisnya
agar mudah dicocokkan.
"""

from __future__ import annotations

import sqlite3
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import config
import db

# --- Peta kolom ----------------------------------------------------------
COL_NO = 1                 # A
COL_TANGGAL = 2            # B (digabung dengan C)
COL_NAMA = 4               # D  ┐
COL_POSISI = 5             # E  │ Orang yang dijamu
COL_PERUSAHAAN = 6         # F  │
COL_INDUSTRI = 7           # G  ┘
COL_LOKASI = 8             # H  ┐
COL_ALAMAT = 9             # I  │ Menjamu klien/mitra bisnis
COL_TIPE = 10              # J  ┘
COL_TUJUAN = 11            # K
COL_JUMLAH = 12            # L
COL_KARYAWAN = 13          # M
LAST_COL = COL_KARYAWAN

COLUMN_WIDTHS = {
    1: 5, 2: 11, 3: 3, 4: 24, 5: 13, 6: 16, 7: 14,
    8: 18, 9: 16, 10: 9, 11: 22, 12: 14, 13: 24,
}

ROW_HEADER_START = 7       # blok identitas karyawan
ROW_TABLE_GROUP = 12       # baris judul gabungan
ROW_TABLE_HEAD = 13        # baris nama kolom
ROW_DATA_START = 14

# --- Gaya ----------------------------------------------------------------
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True, name="Arial Narrow", size=9)
NORMAL = Font(name="Arial Narrow", size=9)
RED_BOLD = Font(bold=True, color="C00000", name="Arial Narrow", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

MONEY_FMT = "#,##0"

CATATAN = ("setelah karyawan selesai melakukan perjalanan dinas dengan "
           "melampirkan semua bukti transaksi/pembayaran asli")

# Nama bulan hidup di web/periods.py supaya bot dan dashboard Vercel memakai
# definisi yang sama, bukan dua salinan yang bisa melenceng. Di-re-export
# karena receipt_image.py memanggilnya lewat modul ini.
from periods import MONTH_ABBR, period_label  # noqa: F401,E402
from periods import fmt_date as _fmt_date     # noqa: E402


def _label(ws: Worksheet, row: int, label: str, value: str) -> None:
    """Satu baris di blok identitas: LABEL : nilai."""
    ws.cell(row=row, column=2, value=label).font = BOLD
    ws.cell(row=row, column=3, value=":").font = BOLD
    cell = ws.cell(row=row, column=4, value=value)
    cell.font = NORMAL
    cell.alignment = LEFT


def _build_sheet(wb: Workbook, sheet_name: str, rows: list[sqlite3.Row], period: str) -> float:
    ws = wb.create_sheet(title=sheet_name)

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    # --- Judul -----------------------------------------------------------
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
    title = ws.cell(row=2, column=1, value=f"FORM {sheet_name.upper()}")
    title.font = Font(bold=True, size=14, name="Arial Narrow")
    title.alignment = CENTER

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=LAST_COL)
    sub = ws.cell(row=3, column=1, value=f"Periode {period_label(period)}")
    sub.font = Font(size=11, name="Arial Narrow")
    sub.alignment = CENTER

    # --- Identitas karyawan ---------------------------------------------
    _label(ws, ROW_HEADER_START + 0, "NPK", config.KARYAWAN_NPK)
    _label(ws, ROW_HEADER_START + 1, "JABATAN", config.KARYAWAN_JABATAN)
    _label(ws, ROW_HEADER_START + 2, "DEPARTEMEN", config.KARYAWAN_DEPARTEMEN)
    _label(ws, ROW_HEADER_START + 3, "PERUSAHAAN", config.PERUSAHAAN)

    # Catatan kecil di kanan atas, seperti di form asli.
    ws.merge_cells(start_row=8, start_column=COL_TUJUAN, end_row=10, end_column=LAST_COL)
    note = ws.cell(row=8, column=COL_TUJUAN, value=CATATAN)
    note.font = Font(italic=True, size=9, name="Arial Narrow")
    note.alignment = CENTER

    # --- Kepala tabel ----------------------------------------------------
    def head(col: int, value: str, span: int = 1, rowspan: int = 1) -> None:
        r = ROW_TABLE_GROUP if rowspan == 2 else ROW_TABLE_HEAD
        if rowspan == 2:
            ws.merge_cells(start_row=ROW_TABLE_GROUP, start_column=col,
                           end_row=ROW_TABLE_HEAD, end_column=col + span - 1)
        elif span > 1:
            ws.merge_cells(start_row=r, start_column=col,
                           end_row=r, end_column=col + span - 1)
        cell = ws.cell(row=r, column=col, value=value)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.fill = HEAD_FILL

    head(COL_NO, "NO", rowspan=2)
    head(COL_TANGGAL, "Tanggal", span=2, rowspan=2)
    head(COL_TUJUAN, "Tujuan", rowspan=2)
    head(COL_JUMLAH, "Jumlah", rowspan=2)
    head(COL_KARYAWAN, "Karyawan Internal", rowspan=2)

    # Dua judul gabungan di baris atas
    ws.merge_cells(start_row=ROW_TABLE_GROUP, start_column=COL_NAMA,
                   end_row=ROW_TABLE_GROUP, end_column=COL_INDUSTRI)
    grp1 = ws.cell(row=ROW_TABLE_GROUP, column=COL_NAMA, value="Orang yang dijamu")
    grp1.font = BOLD
    grp1.alignment = CENTER
    grp1.fill = HEAD_FILL

    ws.merge_cells(start_row=ROW_TABLE_GROUP, start_column=COL_LOKASI,
                   end_row=ROW_TABLE_GROUP, end_column=COL_TIPE)
    grp2 = ws.cell(row=ROW_TABLE_GROUP, column=COL_LOKASI, value="Menjamu klien/mitra bisnis")
    grp2.font = BOLD
    grp2.alignment = CENTER
    grp2.fill = HEAD_FILL

    for col, label in (
        (COL_NAMA, "Nama"), (COL_POSISI, "Posisi"), (COL_PERUSAHAAN, "Perusahaan"),
        (COL_INDUSTRI, "Industri"), (COL_LOKASI, "Lokasi"), (COL_ALAMAT, "Alamat"),
        (COL_TIPE, "Tipe"),
    ):
        head(col, label)

    ws.row_dimensions[ROW_TABLE_GROUP].height = 20
    ws.row_dimensions[ROW_TABLE_HEAD].height = 26

    for r in (ROW_TABLE_GROUP, ROW_TABLE_HEAD):
        for c in range(1, LAST_COL + 1):
            ws.cell(row=r, column=c).border = BORDER

    # --- Isi tabel -------------------------------------------------------
    # Form tetap mencetak minimal 20 baris kosong seperti template aslinya.
    n_rows = max(config.FORM_MIN_ROWS, len(rows))
    total = 0.0

    for i in range(n_rows):
        r = ROW_DATA_START + i
        row = rows[i] if i < len(rows) else None

        ws.merge_cells(start_row=r, start_column=COL_TANGGAL, end_row=r, end_column=3)

        no_cell = ws.cell(row=r, column=COL_NO, value=i + 1)
        no_cell.font = NORMAL
        no_cell.alignment = CENTER

        if row is not None:
            values = {
                COL_TANGGAL: _fmt_date(row["txn_date"]),
                COL_NAMA: row["tamu_nama"],
                COL_POSISI: row["tamu_posisi"],
                COL_PERUSAHAAN: row["tamu_perusahaan"],
                COL_INDUSTRI: row["industri"],
                COL_LOKASI: row["merchant"],
                COL_ALAMAT: row["alamat"],
                COL_TIPE: row["tipe"],
                COL_TUJUAN: row["tujuan"],
                COL_JUMLAH: float(row["jumlah"] or 0),
                COL_KARYAWAN: row["karyawan_internal"],
            }
            total += float(row["jumlah"] or 0)
        else:
            values = {}

        for col in range(COL_TANGGAL, LAST_COL + 1):
            if col == 3:
                continue  # ikut merge dengan kolom Tanggal
            cell = ws.cell(row=r, column=col, value=values.get(col))
            cell.font = NORMAL
            cell.border = BORDER
            if col == COL_JUMLAH:
                cell.number_format = MONEY_FMT
                cell.alignment = RIGHT
            elif col in (COL_TANGGAL, COL_TIPE):
                cell.alignment = CENTER
            else:
                cell.alignment = LEFT

        ws.cell(row=r, column=COL_NO).border = BORDER
        ws.cell(row=r, column=3).border = BORDER
        ws.row_dimensions[r].height = 21

    # --- Sub total -------------------------------------------------------
    data_end = ROW_DATA_START + n_rows - 1
    subtotal_row = data_end + 2

    ws.merge_cells(start_row=subtotal_row, start_column=COL_NO,
                   end_row=subtotal_row, end_column=COL_TUJUAN)
    st_label = ws.cell(row=subtotal_row, column=COL_NO, value="SUB TOTAL")
    st_label.font = BOLD
    st_label.alignment = Alignment(horizontal="center", vertical="center")

    st_value = ws.cell(row=subtotal_row, column=COL_JUMLAH, value=total)
    st_value.font = BOLD
    st_value.number_format = MONEY_FMT
    st_value.alignment = RIGHT
    st_value.border = BORDER

    # --- Grand total / pembayaran dimuka / saldo akhir -------------------
    def total_line(row: int, label: str, value: float | None) -> None:
        ws.merge_cells(start_row=row, start_column=COL_ALAMAT,
                       end_row=row, end_column=COL_TUJUAN)
        lbl = ws.cell(row=row, column=COL_ALAMAT, value=label)
        lbl.font = RED_BOLD
        lbl.alignment = RIGHT

        cur = ws.cell(row=row, column=COL_JUMLAH, value=config.CURRENCY.replace("IDR", "Rp"))
        cur.font = RED_BOLD
        cur.alignment = Alignment(horizontal="left", vertical="center")

        amt = ws.cell(row=row, column=COL_KARYAWAN, value=value)
        amt.font = RED_BOLD
        amt.number_format = MONEY_FMT
        amt.alignment = RIGHT

    grand_row = subtotal_row + 2
    total_line(grand_row, "GRAND TOTAL", total)
    total_line(grand_row + 1, "PEMBAYARAN DIMUKA", 0)
    total_line(grand_row + 2, "SALDO AKHIR", total)

    # --- Blok tanda tangan ----------------------------------------------
    sign_row = grand_row + 5
    blocks = [
        (COL_TANGGAL, "Diajukan oleh,", "Karyawan"),
        (COL_LOKASI, "Disetujui oleh,", "Atasan"),
        (COL_TUJUAN, "Diketahui oleh,", "HRGA Senior Manager"),
    ]
    for col, line1, line2 in blocks:
        ws.cell(row=sign_row, column=col, value=line1).font = NORMAL
        ws.cell(row=sign_row + 1, column=col, value=line2).font = NORMAL
        ws.cell(row=sign_row + 5, column=col, value="Nama    :").font = NORMAL
        ws.cell(row=sign_row + 6, column=col, value="Tanggal :").font = NORMAL

    ws.print_area = f"A1:{get_column_letter(LAST_COL)}{sign_row + 7}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return total


def report_filename(period: str) -> str:
    year, month = period.split("-")
    nama = f" {config.KARYAWAN_NAMA}" if config.KARYAWAN_NAMA else ""
    return f"Reimbursement{nama} {MONTH_ABBR[int(month) - 1]} {year}.xlsx"


def build_workbook(period: str) -> Path:
    """Bangun ulang form Excel untuk satu periode. Selalu menimpa file lama."""
    wb = Workbook()
    wb.remove(wb.active)

    for payment_type, sheet_name in config.SHEET_NAMES.items():
        _build_sheet(wb, sheet_name, db.list_by_period(period, payment_type), period)

    wb.properties.title = f"Form Reimbursement {period_label(period)}"
    wb.properties.created = datetime.now()

    out_path = config.EXPORT_DIR / report_filename(period)
    wb.save(out_path)
    return out_path


def build_photo_archive(period: str) -> Path | None:
    """Kumpulkan foto struk satu periode ke dalam .zip.

    Nama file di dalam zip diawali nomor baris, jadi baris 1 di sheet
    Reimbursement bersesuaian dengan `Reimbursement/001_...`.
    """
    out_path = config.EXPORT_DIR / f"Struk_{period}.zip"
    written = 0

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for payment_type, sheet_name in config.SHEET_NAMES.items():
            for idx, row in enumerate(db.list_by_period(period, payment_type), start=1):
                photo = Path(row["photo_path"])
                if not photo.exists():
                    continue
                safe = "".join(
                    c for c in row["merchant"] if c.isalnum() or c in " -_"
                ).strip()[:40] or "struk"
                zf.write(photo, f"{sheet_name}/{idx:03d}_{row['txn_date']}_{safe}{photo.suffix}")
                written += 1

    if written == 0:
        out_path.unlink(missing_ok=True)
        return None
    return out_path


def build_all(period: str) -> tuple[Path, Path | None]:
    return build_workbook(period), build_photo_archive(period)
